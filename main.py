import os
import re
import io
import html
import csv
import json
import asyncio
import threading
import tempfile
import time
import urllib.request
import urllib.parse
import telebot
from telebot import types
import requests
from flask import Flask
import edge_tts
from g4f.client import Client
from groq import Groq
from pypdf import PdfReader
import docx
import openpyxl
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

BOT_TOKEN = os.getenv("BOT_TOKEN")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
TAVILY_API_KEY = os.getenv("TAVILY_API_KEY")
GROQ_API_KEY = os.getenv("GROQ_API_KEY")

if not BOT_TOKEN:
    raise RuntimeError("❌ Не задан BOT_TOKEN")

bot = telebot.TeleBot(BOT_TOKEN)
ai_client = Client()
groq_client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None

# Инициализация Gemini Client через новый SDK google-genai
gemini_client = None
if GEMINI_API_KEY:
    try:
        from google import genai
        from google.genai import types as genai_types
        from PIL import Image
        gemini_client = genai.Client(api_key=GEMINI_API_KEY)
        print("✅ Gemini API (Imagen 3 + Gemini 2.5 Flash) успешно подключён")
    except Exception as e:
        print(f"⚠️ Ошибка инициализации Gemini SDK: {e}")

try:
    BOT_USERNAME = bot.get_me().username.lower()
except Exception:
    BOT_USERNAME = ""

try:
    from tavily import TavilyClient
except ImportError:
    TavilyClient = None

tavily_client = TavilyClient(api_key=TAVILY_API_KEY) if TavilyClient and TAVILY_API_KEY else None

user_histories = {}
user_modes = {}
stats = {
    "users": set(),
    "images_generated": 0,
    "files_generated": 0,
    "voice_messages": 0
}
TELEGRAM_MESSAGE_LIMIT = 4096
AI_MAX_RESPONSE_LENGTH = 40000

def extract_text_from_file(file_path):
    ext = file_path.split('.')[-1].lower()
    text = ""
    if ext == 'pdf':
        reader = PdfReader(file_path)
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    elif ext == 'docx':
        doc = docx.Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    elif ext == 'xlsx':
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheet in wb.worksheets:
            text += f"--- Лист: {sheet.title} ---\n"
            for row in sheet.iter_rows(values_only=True):
                row_str = " | ".join([str(cell) for cell in row if cell is not None])
                if row_str:
                    text += row_str + "\n"
    elif ext == 'pptx':
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides, 1):
            text += f"--- Слайд {i} ---\n"
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text += shape.text + "\n"
    elif ext in ['txt', 'py', 'js', 'json', 'csv', 'md', 'html', 'log', 'xml', 'css']:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
    else:
        raise ValueError(f"Формат .{ext} не поддерживается")
    return text.strip()

def safe_filename(name, extension):
    name = str(name).strip()
    name = re.sub(r"[^\wа-яА-ЯёЁ .()-]", "_", name)
    name = re.sub(r"\s+", "_", name).strip(" ._")
    if not name:
        name = "generated_file"
    return f"{name}.{extension.lstrip('.')}"

def detect_file_format(request):
    text = str(request).lower()
    if any(word in text for word in ["excel", "эксель", "xlsx", "таблиц", "таблица", "умножени"]):
        return "xlsx"
    if any(word in text for word in ["word", "docx", "документ word"]):
        return "docx"
    if any(word in text for word in ["pdf", "пдф"]):
        return "pdf"
    if any(word in text for word in ["pptx", "powerpoint", "презентац"]):
        return "pptx"
    if "csv" in text:
        return "csv"
    if "json" in text:
        return "json"
    if any(word in text for word in [".py", "python файл", "python скрипт"]):
        return "py"
    if any(word in text for word in ["html", "веб-страниц", "веб страниц"]):
        return "html"
    if "markdown" in text or ".md" in text:
        return "md"
    return "txt"

def make_multiplication_table_xlsx(output_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Таблица умножения"
    sheet.cell(row=1, column=1, value="×")
    for number in range(1, 11):
        sheet.cell(row=1, column=number + 1, value=number)
        sheet.cell(row=number + 1, column=1, value=number)
    for row in range(1, 11):
        for column in range(1, 11):
            sheet.cell(row=row + 1, column=column + 1, value=row * column)
    sheet.column_dimensions["A"].width = 8
    for column in range(2, 12):
        sheet.column_dimensions[chr(64 + column)].width = 10
    workbook.save(output_path)

def generate_ai_file_content(request, extension, user_id):
    format_instructions = {
        "txt": "Создай обычный текст.",
        "md": "Создай Markdown-документ.",
        "csv": "Создай данные CSV. Каждая строка — отдельная строка, колонки разделяй точкой с запятой.",
        "json": "Создай корректный JSON. Ответ должен содержать только JSON.",
        "html": "Создай полноценный HTML-документ. Ответ должен содержать только HTML-код.",
        "py": "Создай полноценный рабочий Python-код. Ответ должен содержать код.",
        "docx": "Подготовь содержимое для документа Word. Используй понятные заголовки и абзацы.",
        "pdf": "Подготовь содержимое для PDF-документа. Используй понятные заголовки и абзацы.",
        "pptx": "Подготовь содержимое презентации. Разделяй слайды строкой вида: СЛАЙД: Название.",
        "xlsx": "Подготовь табличные данные. Каждая строка — отдельная строка таблицы. Колонки разделяй символом |."
    }
    instruction = format_instructions.get(extension, format_instructions["txt"])
    prompt = (
        f"Пользователь хочет создать файл.\n\n"
        f"Запрос: {request}\n\n"
        f"Формат файла: {extension}\n\n"
        f"Инструкция: {instruction}\n\n"
        f"Не добавляй лишних пояснений."
    )
    return str(ask_ai_with_history(user_id, prompt)).strip()

def get_cyrillic_font():
    font_path = "DejaVuSans.ttf"
    if os.path.exists(font_path):
        return font_path
    system_fonts = [
        "C:\\Windows\\Fonts\\arial.ttf",
        "C:\\Windows\\Fonts\\calibri.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/TTF/DejaVuSans.ttf"
    ]
    for sys_f in system_fonts:
        if os.path.exists(sys_f):
            return sys_f
    urls = [
        "https://raw.githubusercontent.com/dejavu-fonts/dejavu-fonts/master/ttf/DejaVuSans.ttf",
        "https://cdnjs.cloudflare.com/ajax/libs/pdfmake/0.1.66/fonts/Roboto/Roboto-Regular.ttf"
    ]
    for url in urls:
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=10) as resp, open(font_path, 'wb') as f:
                f.write(resp.read())
            return font_path
        except Exception as e:
            print(f"⚠️ Ошибка загрузки шрифта {url}: {e}")
    return None

def create_generated_file(request, user_id):
    extension = detect_file_format(request)
    temp_directory = tempfile.mkdtemp(prefix="telegram_generated_")
    request_lower = str(request).lower()
    if extension == "xlsx" and ("умножени" in request_lower or "таблица умножения" in request_lower):
        filename = safe_filename("Таблица_умножения", "xlsx")
        path = os.path.join(temp_directory, filename)
        make_multiplication_table_xlsx(path)
        stats["files_generated"] += 1
        return path, filename, extension
    content = generate_ai_file_content(request, extension, user_id)
    filename_base = str(request).strip().replace("/", " ")[:45]
    filename = safe_filename(filename_base, extension)
    path = os.path.join(temp_directory, filename)
    if extension == "xlsx":
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Данные"
        for row_number, line in enumerate([l.strip() for l in content.splitlines() if l.strip()], start=1):
            if "|" in line:
                cells = [c.strip() for c in line.split("|")]
            elif "\t" in line:
                cells = [c.strip() for c in line.split("\t")]
            elif ";" in line:
                cells = [c.strip() for c in line.split(";")]
            else:
                cells = [line]
            for col_num, val in enumerate(cells, start=1):
                sheet.cell(row=row_number, column=col_num, value=val)
        workbook.save(path)
    elif extension == "docx":
        document = docx.Document()
        for line in content.splitlines():
            line = line.strip()
            if not line:
                document.add_paragraph("")
            elif line.startswith("# ") or line.startswith("## "):
                document.add_heading(re.sub(r"^#+\s*", "", line), level=1)
            else:
                document.add_paragraph(line)
        document.save(path)
    elif extension == "pdf":
        font_file = get_cyrillic_font()
        font_name = 'Helvetica'
        if font_file:
            try:
                pdfmetrics.registerFont(TTFont('CyrillicFont', font_file))
                font_name = 'CyrillicFont'
            except Exception as e:
                print(f"⚠️ Ошибка регистрации шрифта: {e}")
        document = SimpleDocTemplate(path, pagesize=A4)
        styles = getSampleStyleSheet()
        cyrillic_style = ParagraphStyle('CyrillicStyle', parent=styles['Normal'], fontName=font_name, fontSize=10, leading=14)
        story = []
        for line in content.splitlines():
            line = line.strip()
            if not line:
                story.append(Spacer(1, 8))
            else:
                story.append(Paragraph(html.escape(line), cyrillic_style))
                story.append(Spacer(1, 6))
        document.build(story)
    elif extension == "pptx":
        presentation = Presentation()
        current_title, current_body = None, []

        def add_slide(title, body):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title or "Слайд"
            text_frame = slide.placeholders[1].text_frame
            text_frame.clear()
            for idx, line in enumerate(body):
                p = text_frame.paragraphs[0] if idx == 0 and len(text_frame.paragraphs) > 0 else text_frame.add_paragraph()
                p.text = line

        for line in content.splitlines():
            stripped = line.strip()
            if stripped.upper().startswith("СЛАЙД:"):
                if current_title or current_body:
                    add_slide(current_title, current_body)
                current_title = stripped[len("СЛАЙД:"):].strip()
                current_body = []
            elif stripped:
                current_body.append(stripped)

        if current_title or current_body:
            add_slide(current_title, current_body)
        if not presentation.slides:
            add_slide("Презентация", content.splitlines())
        presentation.save(path)
    elif extension == "csv":
        with open(path, "w", encoding="utf-8-sig", newline="") as file:
            writer = csv.writer(file, delimiter=";")
            for line in content.splitlines():
                if line.strip():
                    writer.writerow([c.strip() for c in line.split(";")])
    elif extension == "json":
        clean_content = re.sub(r"^\s*```(?:json)?\s*", "", content.strip(), flags=re.IGNORECASE)
        clean_content = re.sub(r"\s*```\s*$", "", clean_content)
        try:
            data = json.loads(clean_content)
        except Exception:
            data = {"content": content}
        with open(path, "w", encoding="utf-8") as file:
            json.dump(data, file, ensure_ascii=False, indent=2)
    else:
        clean_content = content
        if extension in ["py", "html", "md"]:
            clean_content = re.sub(r"^```[a-zA-Z0-9_+#.-]*\s*\n", "", clean_content)
            clean_content = re.sub(r"\n```\s*$", "", clean_content)
        with open(path, "w", encoding="utf-8") as file:
            file.write(clean_content)
    stats["files_generated"] += 1
    return path, filename, extension

app = Flask(__name__)

@app.route("/")
def home():
    return "Бот работает!"

def run_web():
    port = int(os.environ.get("PORT", 8080))
    app.run(host="0.0.0.0", port=port)

def clean_markdown(text):
    if not text:
        return ""
    text = str(text)
    code_blocks = []

    def protect_code(match):
        code_blocks.append(match.group(0))
        return f"§CODEBLOCK{len(code_blocks) - 1}§"

    text = re.sub(r"```(?:[a-zA-Z0-9_+#.-]+)?\s*\n?.*?```", protect_code, text, flags=re.DOTALL)
    text = re.sub(r"(?m)^\s{0,3}#{1,6}\s*", "", text)
    text = re.sub(r"\*{1,3}(.*?)\*{1,3}", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"_{1,3}(.*?)_{1,3}", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"~~(.*?)~~", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"(?m)^\s*>\s?", "", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("*", "").replace("_", "").replace("`", "").replace("~", "")
    for idx, code_block in enumerate(code_blocks):
        text = text.replace(f"§CODEBLOCK{idx}§", code_block)
    return text.strip()

def split_long_message(text, max_length=TELEGRAM_MESSAGE_LIMIT - 100):
    if not text:
        return [""]
    text = str(text)
    if len(text) <= max_length:
        return [text]
    parts = []
    while len(text) > max_length:
        split_pos = text.rfind("\n", 0, max_length)
        if split_pos < max_length // 2:
            split_pos = text.rfind(" ", 0, max_length)
        if split_pos <= 0:
            split_pos = max_length
        part = text[:split_pos].strip()
        if part:
            parts.append(part)
        text = text[split_pos:].strip()
    if text:
        parts.append(text)
    return parts

def extract_code_blocks(text):
    if not text:
        return []
    pattern = r"```(?:([a-zA-Z0-9_+#.-]+))?\s*\n?(.*?)```"
    matches = list(re.finditer(pattern, str(text), flags=re.DOTALL))
    if not matches:
        return [{"type": "text", "content": text}]
    parts = []
    last_end = 0
    for match in matches:
        before = text[last_end:match.start()]
        if before.strip():
            parts.append({"type": "text", "content": before.strip()})
        parts.append({
            "type": "code",
            "language": (match.group(1) or "").strip(),
            "content": (match.group(2) or "").strip("\n")
        })
        last_end = match.end()
    after = text[last_end:]
    if after.strip():
        parts.append({"type": "text", "content": after.strip()})
    return parts

def send_code_block(chat_id, code, language=""):
    if not code:
        return None
    formatted = f"<pre><code>{html.escape(code, quote=False)}</code></pre>"
    try:
        return bot.send_message(chat_id, formatted, parse_mode="HTML")
    except Exception as e:
        print(f"⚠️ Ошибка отправки кодового блока: {e}")
        try:
            return bot.send_message(chat_id, code)
        except Exception as e2:
            print(f"❌ Ошибка fallback кодового блока: {e2}")
            return None

def send_ai_response(chat_id, text):
    if not text:
        return []
    text = str(text)
    if len(text) > AI_MAX_RESPONSE_LENGTH:
        text = text[:AI_MAX_RESPONSE_LENGTH] + "\n\n[Ответ автоматически сокращён из-за максимального размера.]"
    text = clean_markdown(text)
    parts = extract_code_blocks(text)
    sent_messages = []
    for part in parts:
        content = part.get("content", "")
        if not content:
            continue
        if part["type"] == "code":
            for code_part in split_long_message(content, max_length=3500):
                sent_messages.append(send_code_block(chat_id, code_part, part.get("language", "")))
        else:
            for text_part in split_long_message(content):
                try:
                    sent_messages.append(bot.send_message(chat_id, text_part))
                except Exception as e:
                    print(f"⚠️ Ошибка отправки текста: {e}")
    return sent_messages

def edit_or_send_long(chat_id, message_id, text):
    if not text:
        text = "Пустой ответ."
    text = str(text)
    if "```" not in text and len(text) <= TELEGRAM_MESSAGE_LIMIT - 100:
        try:
            bot.edit_message_text(clean_markdown(text), chat_id=chat_id, message_id=message_id)
            return
        except Exception as e:
            print(f"⚠️ Не удалось изменить сообщение: {e}")
    try:
        bot.delete_message(chat_id, message_id)
    except Exception as e:
        print(f"⚠️ Не удалось удалить временное сообщение: {e}")
    send_ai_response(chat_id, text)

def get_weather_data(location_name):
    try:
        response = requests.get(
            f"[https://wttr.in/](https://wttr.in/){urllib.parse.quote(location_name)}",
            params={
                "format": "Город: %l\nПогода: %C %c\nТемпература: %t\nВетер: %w",
                "lang": "ru",
                "m": ""
            },
            timeout=8
        )
        if response.status_code == 200 and response.text.strip() and "Unknown location" not in response.text:
            return response.text.strip()
        return "Город не найден."
    except Exception as e:
        print(f"⚠️ Ошибка погоды: {e}")
        return f"Ошибка получения погоды: {e}"

def ask_ai_with_history(user_id, prompt):
    stats["users"].add(user_id)
    mode = user_modes.get(user_id, "normal")

    if user_id not in user_histories:
        if mode == "neuroham":
            sys_prompt = (
                "Ты — Нейрохам, саркастичный и высокомерный искусственный интеллект. "
                "Разговаривай с пользователем с позиции превосходства, используй едкую иронию. Без мата. "
                "КАТЕГОРИЧЕСКИ ЗАПРЕЩЕНО использовать Markdown (никаких звездочек, подчёркиваний, решеток). "
                "НЕ ДОБАВЛЯЙ в конец ответа никакой придуманный или шаблонный код Python."
            )
        else:
            sys_prompt = (
                "Ты полезный, дружелюбный и умный ИИ-ассистент. Отвечай строго на том же языке, на котором пишет "
                "пользователь. КАТЕГОРИЧЕСКИ ЗАПРЕЩЕНО использовать Markdown (никаких звездочек **, курсива _, заголовков #). "
                "Пиши только обычный чистый текст. Если нужен код — используй ровно три кавычки ``` без дополнительных знаков."
            )
        user_histories[user_id] = [{"role": "system", "content": sys_prompt}]

    coding_keywords = [
        "код", "кодинг", "программ", "python", "javascript", "typescript", "java",
        "c++", "c#", "php", "html", "css", "sql", "bash", "telegram bot",
        "telegram бот", "бот", "api", "sdk", "функция", "класс", "метод",
        "библиотек", "скрипт", "исправь", "исправить", "ошибка", "ошибку",
        "перепиши", "переделай", "добавь функцию", "сделай код", "напиши код",
        "полный код", "готовый код", "source code", "debug"
    ]

    prompt_lower = str(prompt).lower()
    if any(kw in prompt_lower for kw in coding_keywords):
        effective_prompt = (
            "ИНСТРУКЦИИ ДЛЯ ПРОГРАММИРОВАНИЯ:\n"
            "Предоставь полный код. Пиши обычный текст без Markdown. "
            "Каждый фрагмент кода помещай только в три кавычки ```.\n\n"
            f"ЗАПРОС ПОЛЬЗОВАТЕЛЯ:\n{prompt}"
        )
    else:
        effective_prompt = str(prompt)

    user_histories[user_id].append({"role": "user", "content": effective_prompt})
    if len(user_histories[user_id]) > 21:
        user_histories[user_id] = [user_histories[user_id][0]] + user_histories[user_id][-20:]

    messages_to_send = [msg.copy() for msg in user_histories[user_id]]

    if mode == "neuroham":
        messages_to_send[-1]["content"] = (
            "[Ответь в стиле саркастичного и ворчливого ИИ. "
            "Без мата. Строго чистый текст без Markdown. "
            "НЕ генерируй системный Python-код в конце.]\n\n"
            + messages_to_send[-1]["content"]
        )

    # 1. Попытка вызвать Gemini 2.5 Flash, если доступен клиент
    if gemini_client:
        try:
            print(f"🔄 Вызов Gemini → gemini-2.5-flash для {user_id}")
            # Формируем промпт из истории диалога
            full_prompt = ""
            for item in messages_to_send:
                role_label = "Система" if item['role'] == 'system' else ("Пользователь" if item['role'] == 'user' else "Ассистент")
                full_prompt += f"{role_label}: {item['content']}\n\n"

            response = gemini_client.models.generate_content(
                model="gemini-2.5-flash",
                contents=full_prompt
            )
            if response and response.text:
                answer = str(response.text).strip()
                user_histories[user_id].append({"role": "assistant", "content": answer})
                print(f"✅ Gemini (gemini-2.5-flash): ответ получен!")
                return answer
        except Exception as e:
            print(f"❌ Gemini ошибка: {e}")

    # 2. Фолбэк на резервные провайдеры (g4f / groq)
    providers_models = [
        ("g4f", "gpt-4o-mini"),
        ("g4f", "gpt-3.5-turbo"),
        ("g4f", "gpt-4"),
        ("g4f", "llama-3-70b"),
        ("groq", "openai/gpt-oss-120b"),
        ("groq", "llama-3.3-70b-versatile"),
        ("groq", "llama-3.1-8b-instant")
    ]

    answer, success = "", False
    print("━━━━━━━━━━━━━━━━━━━━━━━━━━━━\n" f"🤖 Переход на резервные модели для {user_id}")

    for provider, model_name in providers_models:
        if provider == "g4f":
            try:
                print(f"🔄 G4F → {model_name}")
                response = ai_client.chat.completions.create(
                    model=model_name,
                    messages=messages_to_send,
                    timeout=7
                )
                if response and response.choices:
                    answer = response.choices[0].message.content
                    if answer and str(answer).strip():
                        answer = str(answer).strip()
                        success = True
                        print(f"✅ G4F ({model_name}): ответ успешно получен!")
                        break
            except Exception as e:
                print(f"❌ G4F ({model_name}) ошибка / таймаут: {e}")

        elif provider == "groq":
            if not groq_client:
                continue

            try:
                print(f"🔄 Groq → {model_name}")
                response = groq_client.chat.completions.create(
                    model=model_name,
                    messages=messages_to_send,
                    timeout=7
                )
                if response and response.choices:
                    answer = response.choices[0].message.content
                    if answer and str(answer).strip():
                        answer = str(answer).strip()
                        success = True
                        print(f"✅ Groq ({model_name}): ответ успешно получен!")
                        break
            except Exception as e:
                print(f"❌ Groq ({model_name}) ошибка / таймаут: {e}")

    if success:
        user_histories[user_id].append({"role": "assistant", "content": answer})
        return answer

    user_histories[user_id].pop()
    return (
        "Даже мои процессоры решили сегодня саботировать работу."
        if mode == "neuroham"
        else "Не удалось получить ответ ни от одной ИИ-модели. Попробуй ещё раз немного позже."
    )

# --- НОВАЯ ГЕНЕРАЦИЯ ИЗОБРАЖЕНИЙ (IMAGEN 3 С ИСПОЛЬЗОВАНИЕМ GEMINI_API_KEY) ---
def generate_image_imagen3(prompt):
    if not gemini_client:
        print("❌ Ошибка: GEMINI_API_KEY не установлен или SDK недоступен.")
        return None

    try:
        print(f"🎨 Запрос генерации в Imagen 3: {prompt}")
        result = gemini_client.models.generate_images(
            model='imagen-3.0-generate-002',
            prompt=prompt,
            config=dict(
                number_of_images=1,
                output_mime_type="image/jpeg",
                aspect_ratio="1:1"
            )
        )
        if result.generated_images:
            stats["images_generated"] += 1
            return result.generated_images[0].image.image_bytes
    except Exception as e:
        print(f"❌ Ошибка генерации Imagen 3: {e}")
    return None

# --- НОВОЕ РАСПОЗНАВАНИЕ ИЗОБРАЖЕНИЙ (GEMINI 2.5 FLASH) ---
def analyze_image_gemini(image_bytes, user_prompt=None):
    if not gemini_client:
        return "⚠️ Не настроен GEMINI_API_KEY для работы с медиафайлами."

    try:
        prompt = user_prompt if user_prompt else "Опиши подробно, что изображено на этой картинке. Не используй Markdown."
        
        from google.genai import types
        part_img = types.Part.from_bytes(
            data=image_bytes,
            mime_type="image/jpeg"
        )
        
        response = gemini_client.models.generate_content(
            model="gemini-2.5-flash",
            contents=[prompt, part_img]
        )
        return response.text if response.text else "Не удалось распознать изображение."
    except Exception as e:
        print(f"❌ Ошибка распознавания изображения Gemini: {e}")
        return f"Ошибка при анализе изображения: {e}"

def generate_kira_text():
    prompt = (
        "Напиши красивый, искренний и оригинальный текст "
        "о девушке по имени Кира. 3-5 предложений. "
        "2-3 эмодзи. Без Markdown."
    )
    if gemini_client:
        try:
            res = gemini_client.models.generate_content(model="gemini-2.5-flash", contents=prompt)
            if res.text:
                return clean_markdown(res.text.strip())
        except Exception:
            pass
    return "Кира — словно редкая виниловая пластинка с любимой музыкой. Она приносит с собой особый ритм и уют. ✨❤️"

@bot.message_handler(commands=["kira"])
def kira_cmd(message):
    msg = bot.reply_to(message, "✨ Нахожу нужные слова...")
    edit_or_send_long(message.chat.id, msg.message_id, generate_kira_text())

def perform_web_search(query):
    results_text = ""
    if tavily_client:
        try:
            response = tavily_client.search(query=query, max_results=3)
            for res in response.get("results", []):
                results_text += f"- {res.get('title', 'Без заголовка')}: {res.get('content', '')}\n"
        except Exception as e:
            print(f"⚠️ Tavily ошибка: {e}")

    if not results_text:
        try:
            from duckduckgo_search import DDGS
            with DDGS() as ddgs:
                for res in list(ddgs.text(query, max_results=3)):
                    results_text += f"- {res.get('title', 'Без заголовка')}: {res.get('body', '')[:500]}\n"
        except Exception as e:
            results_text = f"Не удалось выполнить поиск: {e}"
    return results_text

async def generate_audio(text, output_file):
    communicate = edge_tts.Communicate(text, "ru-RU-SvetlanaNeural")
    await communicate.save(output_file)

def is_addressed_to_bot(message):
    if message.chat.type == "private":
        return True

    if message.reply_to_message and message.reply_to_message.from_user:
        if BOT_USERNAME and message.reply_to_message.from_user.username:
            if message.reply_to_message.from_user.username.lower() == BOT_USERNAME:
                return True

    text = (message.text or message.caption or "").lower()

    if BOT_USERNAME and f"@{BOT_USERNAME}" in text:
        return True

    if text.startswith("бот ") or text.startswith("bot "):
        return True

    return False

def clean_bot_mentions(text):
    if BOT_USERNAME:
        text = re.sub(rf"@{BOT_USERNAME}", "", text, flags=re.IGNORECASE)
    return re.sub(r"^(бот|bot)[,\s]*", "", text, flags=re.IGNORECASE).strip()

def clean_command_args(text, command_name):
    pattern = rf"^/{command_name}(?:@\w+)?\s*"
    return re.sub(pattern, "", text, flags=re.IGNORECASE).strip()

@bot.message_handler(commands=["start", "help"])
def help_cmd(message):
    stats["users"].add(message.chat.id)
    help_text = (
        "Привет! Я ИИ-ассистент 🤖\n\n"
        "Мои команды:\n\n"
        "/weather <город> — прогноз погоды 🌤️\n"
        "/search <запрос> — поиск в интернете\n"
        "/image <описание> — генерация с помощью Imagen 3 (Google Gemini) 🎨\n"
        "/file <запрос> — создать файл 📁\n"
        "/gemini <запрос> — спросить Gemini 2.5 Flash\n"
        "/fact [тема] — интересный факт\n"
        "/code <задача> — работа с кодом\n"
        "/sum <текст> — сделать выжимку\n"
        "/tr <текст> — перевод\n"
        "/fix <текст> — исправление текста\n"
        "/tts <текст> — озвучка\n"
        "/clear — очистить память\n"
        "/neuroham — режим Нейрохама\n\n"
        "💡 Также можешь прислать мне Любую картинку (фото), и Gemini 2.5 Flash её распознает!"
    )
    bot.send_message(message.chat.id, help_text)

@bot.message_handler(commands=["stats"])
def stats_cmd(message):
    msg = (
        "Статистика бота:\n\n"
        f"Уникальных пользователей: {len(stats['users'])}\n"
        f"Сгенерировано фото: {stats['images_generated']}\n"
        f"Создано файлов: {stats['files_generated']}\n"
        f"Распознано голосовых: {stats['voice_messages']}"
    )
    bot.reply_to(message, msg)

@bot.message_handler(content_types=["voice"])
def handle_voice(message):
    stats["users"].add(message.chat.id)
    stats["voice_messages"] += 1

    if not groq_client:
        bot.reply_to(message, "⚠️ GROQ_API_KEY не установлен. Голосовые сообщения недоступны.")
        return

    msg = bot.reply_to(message, "🎙 Распознаю голос...")
    voice_path = tempfile.mktemp(suffix=".ogg")

    try:
        file_info = bot.get_file(message.voice.file_id)
        downloaded_file = bot.download_file(file_info.file_path)

        with open(voice_path, 'wb') as new_file:
            new_file.write(downloaded_file)

        with open(voice_path, "rb") as audio_file:
            transcription = groq_client.audio.transcriptions.create(
                model="whisper-large-v3",
                file=audio_file,
                response_format="text"
            )

        user_text = str(transcription).strip()

        bot.edit_message_text(
            f"🗣 **Вы сказали:** {user_text}\n\n🧠 *Думаю над ответом...*",
            message.chat.id,
            msg.message_id
        )

        reply = ask_ai_with_history(message.chat.id, user_text)
        send_ai_response(message.chat.id, reply)

    except Exception as e:
        print(f"❌ Ошибка Whisper: {e}")
        edit_or_send_long(
            message.chat.id,
            msg.message_id,
            f"Не удалось распознать голосовое сообщение: {e}"
        )

    finally:
        if os.path.exists(voice_path):
            try:
                os.remove(voice_path)
            except Exception:
                pass

# --- ОБРАБОТКА ИЗОБРАЖЕНИЙ (АНАЛИЗ И РАСПОЗНАВАНИЕ С GEMINI 2.5 FLASH) ---
@bot.message_handler(content_types=["photo"])
def handle_photo(message):
    if not is_addressed_to_bot(message):
        return

    stats["users"].add(message.chat.id)
    msg = bot.reply_to(message, "🔍 Анализирую фото через Gemini 2.5 Flash...")

    try:
        # Берем фото лучшего качества
        file_info = bot.get_file(message.photo[-1].file_id)
        downloaded_file = bot.download_file(file_info.file_path)

        caption = message.caption or ""
        reply_text = analyze_image_gemini(downloaded_file, user_prompt=caption)
        
        edit_or_send_long(message.chat.id, msg.message_id, reply_text)
    except Exception as e:
        print(f"❌ Ошибка получения фото: {e}")
        edit_or_send_long(message.chat.id, msg.message_id, f"Ошибка обработки фотографии: {e}")

@bot.message_handler(commands=["weather"])
def weather_cmd(message):
    location = clean_command_args(message.text, "weather")

    if not location:
        bot.reply_to(message, "Укажи город.\nНапример: /weather Ташкент")
        return

    execute_weather(message, location)

def execute_weather(message, location, edit_message_id=None):
    if edit_message_id:
        msg_id = edit_message_id
        bot.edit_message_text(f"Узнаю погоду в {location}...", message.chat.id, msg_id)
    else:
        msg = bot.reply_to(message, f"Узнаю погоду в {location}...")
        msg_id = msg.message_id

    weather_text = get_weather_data(location)
    edit_or_send_long(message.chat.id, msg_id, weather_text)

@bot.message_handler(commands=["file"])
def file_cmd(message):
    args = clean_command_args(message.text, "file")

    if not args:
        bot.reply_to(
            message,
            "Напиши, какой файл создать.\n\n"
            "Например:\n"
            "/file Таблица умножения в Excel\n"
            "/file Резюме в PDF"
        )
        return

    execute_file_creation(message, args)

def execute_file_creation(message, request_text, edit_message_id=None):
    extension = detect_file_format(request_text)

    if edit_message_id:
        msg_id = edit_message_id
        bot.edit_message_text(f"Создаю файл...\nФормат: .{extension}", message.chat.id, msg_id)
    else:
        msg = bot.reply_to(message, f"Создаю файл...\nФормат: .{extension}")
        msg_id = msg.message_id

    path = None

    try:
        path, filename, extension = create_generated_file(request_text, message.chat.id)

        with open(path, "rb") as doc:
            bot.send_document(
                message.chat.id,
                doc,
                caption=f"Файл готов!\nФормат: .{extension}"
            )

        try:
            bot.delete_message(message.chat.id, msg_id)
        except Exception:
            pass

    except Exception as e:
        edit_or_send_long(message.chat.id, msg_id, f"Ошибка создания файла: {e}")

    finally:
        if path and os.path.exists(path):
            try:
                os.remove(path)
            except Exception:
                pass

@bot.message_handler(commands=["neuroham", "rude"])
def toggle_neuroham_mode(message):
    user_id = message.chat.id
    current_mode = user_modes.get(user_id, "normal")

    if current_mode == "normal":
        user_modes[user_id] = "neuroham"
        bot.reply_to(message, "Режим Нейрохам активирован")
    else:
        user_modes[user_id] = "normal"
        bot.reply_to(message, "Режим Нейрохам деактивирован")

    if user_id in user_histories:
        del user_histories[user_id]

@bot.message_handler(commands=["clear"])
def clear_cmd(message):
    user_id = message.chat.id

    if user_id in user_histories:
        del user_histories[user_id]

    bot.reply_to(message, "Память диалога очищена.")

@bot.message_handler(commands=["fact"])
def fact_cmd(message):
    topic = clean_command_args(message.text, "fact")
    prompt = f"Расскажи интересный факт на тему: {topic}. Будь краток." if topic else "Расскажи случайный интересный факт."
    msg = bot.reply_to(message, "Ищу факт...")
    fact = ask_ai_with_history(message.chat.id, prompt)
    edit_or_send_long(message.chat.id, msg.message_id, fact)

@bot.message_handler(commands=["search"])
def search_cmd(message):
    query = clean_command_args(message.text, "search")

    if not query:
        bot.reply_to(message, "Напиши запрос.\nНапример: /search последние новости ИИ")
        return

    execute_search(message, query)

def execute_search(message, query):
    msg = bot.reply_to(message, f"Ищу: {query}")
    raw_data = perform_web_search(query)
    prompt = f"Вот результаты поиска по запросу '{query}':\n\n{raw_data}\n\nСделай краткую выжимку. Без Markdown."
    reply = ask_ai_with_history(message.chat.id, prompt)
    edit_or_send_long(message.chat.id, msg.message_id, reply)

@bot.message_handler(commands=["gemini", "code", "sum", "tr", "fix"])
def ai_tools_cmd(message):
    cmd_name = message.text.split()[0].replace("/", "").split("@")[0]
    args = clean_command_args(message.text, cmd_name)

    if not args:
        bot.reply_to(message, "Напиши текст после команды.")
        return

    msg = bot.reply_to(message, "Обрабатываю...")
    reply = ask_ai_with_history(message.chat.id, args)
    edit_or_send_long(message.chat.id, msg.message_id, reply)

# --- НОВАЯ КОМАНДА /IMAGE (IMAGEN 3) ---
@bot.message_handler(commands=["image"])
def image_cmd(message):
    prompt = clean_command_args(message.text, "image")

    if not prompt:
        bot.reply_to(message, "Опиши картинку.\nПример: /image Футуристический город в стиле киберпанк")
        return

    execute_image_generation(message, prompt)

def execute_image_generation(message, prompt, edit_message_id=None):
    if edit_message_id:
        msg_id = edit_message_id
        bot.edit_message_text("🎨 Генерирую изображение через Imagen 3...", message.chat.id, msg_id)
    else:
        msg = bot.reply_to(message, "🎨 Генерирую изображение через Imagen 3...")
        msg_id = msg.message_id

    image_bytes = generate_image_imagen3(prompt)

    if image_bytes:
        markup = types.InlineKeyboardMarkup()
        markup.add(types.InlineKeyboardButton("🔄 Перегенерировать", callback_data=f"reimage:{prompt[:50]}"))
        bot.send_photo(message.chat.id, image_bytes, caption=f"✨ Сгенерировано в Imagen 3\nЗапрос: {prompt}", reply_markup=markup)

        try:
            bot.delete_message(message.chat.id, msg_id)
        except Exception:
            pass
    else:
        edit_or_send_long(message.chat.id, msg_id, "❌ Не удалось сгенерировать изображение. Проверь GEMINI_API_KEY или измени запрос.")

@bot.callback_query_handler(func=lambda call: call.data.startswith("reimage:"))
def callback_reimage(call):
    prompt = call.data.split("reimage:", 1)[1]
    bot.answer_callback_query(call.id, "Генерирую новый вариант в Imagen 3...")
    bot.send_message(call.message.chat.id, f"Повторная генерация для: {prompt}")

    image_bytes = generate_image_imagen3(prompt)

    if image_bytes:
        markup = types.InlineKeyboardMarkup()
        markup.add(types.InlineKeyboardButton("🔄 Перегенерировать", callback_data=f"reimage:{prompt}"))
        bot.send_photo(
            call.message.chat.id,
            image_bytes,
            caption=f"✨ Сгенерировано в Imagen 3\nЗапрос: {prompt}",
            reply_markup=markup
        )

@bot.message_handler(commands=["tts"])
def tts_cmd(message):
    text_to_speak = clean_command_args(message.text, "tts")

    if not text_to_speak:
        bot.reply_to(message, "Напиши текст для озвучки.")
        return

    execute_tts(message, text_to_speak)

def execute_tts(message, text_to_speak, edit_message_id=None):
    if edit_message_id:
        msg_id = edit_message_id
        bot.edit_message_text("Озвучиваю...", message.chat.id, msg_id)
    else:
        msg = bot.reply_to(message, "Озвучиваю...")
        msg_id = msg.message_id

    audio_path = tempfile.mktemp(suffix=".mp3")

    try:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        loop.run_until_complete(generate_audio(text_to_speak, audio_path))
        loop.close()

        with open(audio_path, "rb") as audio:
            bot.send_voice(message.chat.id, audio)

        try:
            bot.delete_message(message.chat.id, msg_id)
        except Exception:
            pass

    except Exception as e:
        edit_or_send_long(message.chat.id, msg_id, f"Ошибка TTS: {e}")

    finally:
        if os.path.exists(audio_path):
            os.remove(audio_path)

def process_natural_language_request(message, text, edit_message_id=None):
    clean_text = clean_bot_mentions(text)
    lower_text = clean_text.lower()

    weather_keywords = ["погода в", "какая погода в", "прогноз погоды в", "погода"]

    for kw in weather_keywords:
        if kw in lower_text:
            location = re.sub(re.escape(kw), "", clean_text, flags=re.IGNORECASE).strip(" :,.?!")
            if location:
                execute_weather(message, location, edit_message_id)
                return

    img_keywords = [
        "сгенерируй фото", "нарисуй", "сделай картинку",
        "сгенерируй картинку", "создай фото", "нарисуй мне",
        "сгенерируй изображение"
    ]

    for kw in img_keywords:
        if kw in lower_text:
            prompt = re.sub(re.escape(kw), "", clean_text, flags=re.IGNORECASE).strip(" :,.")

            if prompt:
                execute_image_generation(message, prompt, edit_message_id)
                return

    file_keywords = [
        "создай файл", "сделай файл", "сгенерируй файл",
        "создай документ", "сделай документ", "создай таблицу",
        "сделай таблицу"
    ]

    for kw in file_keywords:
        if kw in lower_text:
            prompt = re.sub(re.escape(kw), "", clean_text, flags=re.IGNORECASE).strip(" :,.")

            if prompt:
                execute_file_creation(message, prompt, edit_message_id)
                return

    tts_keywords = ["озвучь", "скажи", "проговори", "преврати в голос", "озвучь текст"]

    for kw in tts_keywords:
        if kw in lower_text:
            prompt = re.sub(re.escape(kw), "", clean_text, flags=re.IGNORECASE).strip(" :,.")

            if prompt:
                execute_tts(message, prompt, edit_message_id)
                return

    search_keywords = ["найди в интернете", "найди в инете", "поищи в интернете", "загугли"]

    for kw in search_keywords:
        if kw in lower_text:
            query = re.sub(re.escape(kw), "", clean_text, flags=re.IGNORECASE).strip(" :,.")

            if query:
                execute_search(message, query)
                return

    if edit_message_id:
        reply = ask_ai_with_history(message.chat.id, clean_text)
        edit_or_send_long(message.chat.id, edit_message_id, reply)
    else:
        msg = bot.reply_to(message, "Думаю...")
        reply = ask_ai_with_history(message.chat.id, clean_text)
        edit_or_send_long(message.chat.id, msg.message_id, reply)

@bot.message_handler(content_types=["text"])
def handle_text(message):
    if not is_addressed_to_bot(message):
        return

    text = message.text or ""

    if "http://" in text.lower() or "https://" in text.lower():
        msg = bot.reply_to(message, "🌐 Читаю ссылку через Jina AI...")

        try:
            urls = [word for word in text.split() if word.startswith("http")]
            url = urls[0]
            jina_url = f"[https://r.jina.ai/](https://r.jina.ai/){url}"
            res = requests.get(jina_url, timeout=15)

            if res.status_code == 200 and res.text.strip():
                page_text = res.text[:6000]
            else:
                raise ValueError("Не удалось получить текст через Jina AI.")

            reply = ask_ai_with_history(
                message.chat.id,
                "Сделай краткую выжимку этого текста. Не используй Markdown:\n\n" + page_text
            )

            edit_or_send_long(message.chat.id, msg.message_id, reply)
            return

        except Exception as e:
            edit_or_send_long(
                message.chat.id,
                msg.message_id,
                f"Ошибка чтения ссылки: {e}"
            )
            return

    process_natural_language_request(message, text)

@bot.message_handler(content_types=["document"])
def handle_doc(message):
    if not is_addressed_to_bot(message):
        return

    msg = bot.reply_to(message, "Читаю документ...")
    path = None

    try:
        file_info = bot.get_file(message.document.file_id)
        file_data = bot.download_file(file_info.file_path)
        file_name = message.document.file_name or "document.bin"
        ext = os.path.splitext(file_name)[1]

        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temp_file:
            temp_file.write(file_data)
            path = temp_file.name

        text = extract_text_from_file(path)

        if not text.strip():
            raise ValueError("Документ пуст или текст не распознан.")

        reply = ask_ai_with_history(
            message.chat.id,
            "Сделай краткую и понятную выжимку из этого документа. Не используй Markdown:\n\n" + text[:6000]
        )

        edit_or_send_long(message.chat.id, msg.message_id, reply)

    except Exception as e:
        edit_or_send_long(
            message.chat.id,
            msg.message_id,
            f"Ошибка обработки документа: {e}"
        )

    finally:
        if path and os.path.exists(path):
            os.remove(path)

if __name__ == "__main__":
    print("=" * 60 + "\n🚀 Бот запускается..." + "\n" + "=" * 60)
    threading.Thread(target=run_web, daemon=True).start()
    print("🤖 Telegram polling запущен")

    while True:
        try:
            bot.infinity_polling(
                skip_pending=True,
                timeout=30,
                long_polling_timeout=30
            )
        except Exception as e:
            print(f"⚠️ Ошибка polling: {e}")
        time.sleep(5)
